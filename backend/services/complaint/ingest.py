"""客诉工单入库 —— 导出文件归档 → 解析 → 去重入库 → LLM 归类 → 批量簇/复发标记。

用法（在 backend 目录下执行）：
    python -m scripts.complaint.ingest "D:\\path\\小信工单上周全部明细.xlsx"
    python -m scripts.complaint.ingest <file> --dry-run     # 只解析不入库

设计要点（方案 v3.0）：
  - 原始文件按自然周归档到 data/complaint/YYYY-Www/，永久留底可溯源；
  - ticket_no（BOMC流水号）唯一去重，重复导入天然幂等；
  - **不定期上传**：week 由工单创建时间派生（仅作分析视角），与导入频率解耦；
  - 分类走 **LLM 主导**（llm_classify.py，无规则硬匹配）：新工单先落"待分析"
    （category_l1 空），导入完成后后台线程触发 LLM 归类，完成后自动重算
    批量簇/复发与共性问题识别；
  - 人工校正 classify_by='manual' 永不被覆盖；
  - 批量簇/复发为派生字段，归类完成后按涉及周重算，保证口径一致。
"""

from __future__ import annotations

import shutil
import threading
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from services.complaint import config as C


# ── 工具 ──────────────────────────────────────────────────────
def week_key(dt: Optional[datetime]) -> str:
    """自然周（ISO 周）标识，如 2026-W44。"""
    if not dt:
        return ""
    y, w, _ = dt.isocalendar()
    return f"{y}-W{w:02d}"


def _to_dt(v) -> Optional[datetime]:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.replace(tzinfo=None)
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    s = str(v).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S",
                "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def _s(v, limit: int = 255) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return s[:limit]


# ── 解析 ──────────────────────────────────────────────────────
def read_rows(path: Path) -> Tuple[List[Dict], List[str]]:
    """读取导出文件，返回 [{列名: 值}] 与表头列表。"""
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows_iter)]
    out: List[Dict] = []
    for values in rows_iter:
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        out.append(dict(zip(header, values)))
    wb.close()
    return out, header


def archive_source(path: Path, week: str) -> Path:
    """原始导出按自然周归档留底。"""
    target_dir = C.ARCHIVE_ROOT / (week or "unknown")
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / path.name
    if not target.exists():
        shutil.copy2(path, target)
    return target


# ── 入库 ──────────────────────────────────────────────────────
def _bg_llm_classify(weeks: Sequence[str]) -> None:
    """后台线程：LLM 归类 + 派生重算 + 共性问题识别（独立会话，异常全吞记日志）。

    线程内 SessionLocal 延迟导入（避免测试环境误连真实库）。
    """
    import logging

    from db.base import SessionLocal
    from services.complaint.llm_classify import run_llm_classify
    from services.complaint.problems import after_import

    logger = logging.getLogger(__name__)
    try:
        with SessionLocal() as s:
            res = run_llm_classify(s, force=False)
            logger.info("导入后 LLM 归类完成：%s", res)
            ws = [w for w in weeks if w]
            refresh_derived(s, weeks=ws)
            try:
                after_import(s, weeks=ws)
            except Exception:  # noqa: BLE001
                logger.exception("导入后共性问题识别失败")
    except Exception:  # noqa: BLE001
        logger.exception("导入后 LLM 归类线程异常")


def ingest_file(db: Session, path: Path, import_batch: str = "",
                run_llm: bool = True) -> Dict:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)

    rows, header = read_rows(path)
    missing = [c for c in C.COL_HEADER_CANDIDATES if c not in header]
    if missing:
        raise ValueError(f"导出文件缺少必需列：{missing}；实际表头={header}")

    dic = None  # 规则引擎已废弃（LLM 主导），保留变量位防下游误用
    existing = {r[0] for r in db.execute(text("SELECT ticket_no FROM pmwb_complaint_ticket"))}

    batch = import_batch or datetime.now().strftime("%Y%m%d%H%M%S")
    archive_path = None
    inserted = skipped = 0
    weeks = set()

    for row in rows:
        ticket_no = _s(row.get(C.COL_TICKET_NO), 64)
        if not ticket_no:
            skipped += 1
            continue
        if ticket_no in existing:
            skipped += 1
            continue

        created = _to_dt(row.get(C.COL_CREATED))
        week = week_key(created)
        if archive_path is None and week:
            archive_path = archive_source(path, week)
        weeks.add(week)

        content = _s(row.get(C.COL_CONTENT), 20000)
        solution = _s(row.get(C.COL_SOLUTION), 20000)
        # LLM 主导：新工单先落待分析（l1/l2 空），导入完成后后台线程统一归类
        l1, l2 = "", ""

        dispatch_raw = row.get(C.COL_DISPATCH)
        try:
            dispatch = int(float(dispatch_raw)) if dispatch_raw not in (None, "") else 1
        except (TypeError, ValueError):
            dispatch = 1

        db.execute(text("""
            INSERT INTO pmwb_complaint_ticket
                (ticket_no, week, created_time, city, channel, appeal_category, biz_attr,
                 content, solution_history, handler, handle_unit, dispatch_count, status,
                 category_l1, category_l2, classify_by, cluster_key, is_batch, is_recur,
                 source_file, import_batch, created_at, updated_at)
            VALUES
                (:ticket_no, :week, :created_time, :city, :channel, :appeal, :biz_attr,
                 :content, :solution, :handler, :unit, :dispatch, :status,
                 :l1, :l2, 'auto', NULL, 0, 0, :src, :batch, NOW(), NOW())
        """), {
            "ticket_no": ticket_no, "week": week, "created_time": created,
            "city": _s(row.get(C.COL_CITY), 64), "channel": _s(row.get(C.COL_CHANNEL), 64),
            "appeal": _s(row.get(C.COL_APPEAL), 64), "biz_attr": _s(row.get(C.COL_BIZ_ATTR), 64),
            "content": content, "solution": solution,
            "handler": _s(row.get(C.COL_HANDLER), 255), "unit": _s(row.get(C.COL_UNIT), 255),
            "dispatch": dispatch, "status": _s(row.get(C.COL_STATUS), 32),
            "l1": l1, "l2": l2, "src": path.name, "batch": batch,
        })
        existing.add(ticket_no)
        inserted += 1

    db.commit()

    derived = refresh_derived(db, weeks=sorted(w for w in weeks if w))

    # LLM 主导归类：导入成功后后台线程触发（失败不影响导入本身，工单保持待分析）
    llm_thread = None
    if run_llm and inserted > 0:
        llm_thread = threading.Thread(
            target=_bg_llm_classify, args=(sorted(w for w in weeks if w),),
            daemon=True, name="complaint-llm-classify",
        )
        llm_thread.start()

    return {
        "file": path.name,
        "archive": str(archive_path) if archive_path else "",
        "parsed": len(rows),
        "inserted": inserted,
        "skipped_duplicate": skipped,
        "weeks": sorted(w for w in weeks if w),
        "batch": batch,
        "clusters": derived["clusters"],
        "recur": derived["recur"],
        "llm_classify_started": bool(llm_thread),
    }


# ── 派生字段重算 ──────────────────────────────────────────────
def refresh_derived(db: Session, weeks: Optional[Sequence[str]] = None,
                    threshold: int = C.BATCH_DAILY_THRESHOLD) -> Dict:
    """重算批量簇标识与复发标记（整表或指定周）。

    批量簇：同二级场景 + 同自然日单量 >= 阈值 → cluster_key='场景@日期'；
    复发：措施台账该场景已有生效日期，且工单创建时间在其之后 → is_recur=1。
    先清零再打标，避免历史标记残留。

    实现说明：聚合在 Python 侧完成后逐单更新（SQL 方言可移植，
    被 API /correct-category 在测试库 SQLite 上调用；周数据量 ≤ 千级，性能无虞）。
    """
    f_week = ""
    params: Dict = {}
    if weeks:
        for i, w in enumerate(weeks):
            params[f"w{i}"] = w
        in_list = ",".join(f":w{i}" for i in range(len(weeks)))
        f_week = f" AND week IN ({in_list})"

    rows = db.execute(text(
        "SELECT id, category_l2, created_time FROM pmwb_complaint_ticket "
        "WHERE 1 = 1" + f_week
    ), params).fetchall()

    # 1) 批量异常簇：清零 → 按场景+自然日计数打标
    db.execute(text(
        "UPDATE pmwb_complaint_ticket SET cluster_key = NULL, is_batch = 0 WHERE 1 = 1" + f_week
    ), params)

    from collections import Counter
    day_counter: Counter = Counter()
    for _tid, l2, created in rows:
        if l2 and created:
            day_counter[(l2, created.date())] += 1

    cluster_updates = []
    for tid, l2, created in rows:
        if l2 and created and day_counter[(l2, created.date())] >= threshold:
            cluster_updates.append(
                {"id": tid, "ck": f"{l2}@{created.date().isoformat()}"})
    if cluster_updates:
        db.execute(text(
            "UPDATE pmwb_complaint_ticket SET cluster_key = :ck, is_batch = 1 WHERE id = :id"
        ), cluster_updates)

    # 2) 措施上线后同类复发：清零 → 生效措施匹配同场景打标
    db.execute(text(
        "UPDATE pmwb_complaint_ticket SET is_recur = 0 WHERE 1 = 1" + f_week
    ), params)
    go_lives = db.execute(text(
        "SELECT category_l2, go_live_date FROM pmwb_complaint_action "
        "WHERE go_live_date IS NOT NULL "
        "AND status IN ('verifying', 'effective', 'ineffective')"
    )).fetchall()
    recur_map: Dict[str, list] = {}
    for l2, g in go_lives:
        recur_map.setdefault(l2 or "", []).append(g)

    recur_updates = []
    for tid, l2, created in rows:
        if l2 and created:
            for g in recur_map.get(l2, []):
                g_dt = g if isinstance(g, datetime) else datetime.combine(g, datetime.min.time())
                if created >= g_dt:
                    recur_updates.append({"id": tid})
                    break
    if recur_updates:
        db.execute(text(
            "UPDATE pmwb_complaint_ticket SET is_recur = 1 WHERE id = :id"
        ), recur_updates)

    db.commit()

    clusters = sum(1 for k in day_counter.values() if k >= threshold)
    recur = len(recur_updates)
    return {"clusters": int(clusters), "recur": int(recur)}


def reclassify(db: Session) -> Dict:
    """重跑归类（已转 LLM 主导，转发 llm_classify.run_llm_classify）。

    兼容保留旧入口名：force=True 重分析全部非人工校正工单。
    """
    from services.complaint.llm_classify import run_llm_classify
    return run_llm_classify(db, force=True)


